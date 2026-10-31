import json, pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
S=json.load(open("results/summary.json")); km=pd.read_csv("results/kmeans_k_selection.csv"); prof=pd.read_csv("results/kmeans_profiles.csv"); pdb=pd.read_csv("results/dbscan_profiles.csv")
d=Document()
for s in d.sections: s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(2)
d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(11)
def H(t,l=1): d.add_heading(t,level=l)
def P(t,i=False,b=False):
    p=d.add_paragraph(); r=p.add_run(t); r.italic=i; r.bold=b; p.paragraph_format.space_after=Pt(6); return p
def B(items):
    for t in items: d.add_paragraph(t,style="List Bullet")
def FIG(f,cap,w=15.5):
    d.add_picture(f"figures/{f}",width=Cm(w)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(); r=p.add_run(cap); r.italic=True; r.font.size=Pt(9); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def TBL(df,widths=None,fs=9):
    t=d.add_table(rows=1,cols=len(df.columns)); t.style="Light Grid Accent 1"
    for i,c in enumerate(df.columns):
        t.rows[0].cells[i].text=""; r=t.rows[0].cells[i].paragraphs[0].add_run(str(c)); r.bold=True; r.font.size=Pt(fs)
    for _,row in df.iterrows():
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=""; r=cs[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(fs)
    d.add_paragraph()
d.add_heading("Hospital Patient Segmentation using K-Means and DBSCAN",0)
P("Machine Learning Mini Project. Topic 19: Clustering",b=True)
t=d.add_table(rows=1,cols=2); t.style="Light Grid Accent 1"; t.rows[0].cells[0].text="Team member"; t.rows[0].cells[1].text="Register number"
for n,r in [("Tamilmani B","111624102217"),("Yuvraj U","111624102247"),("Vishnu Vardhan A","111624102241")]:
    c=t.add_row().cells; c[0].text=n; c[1].text=r
P("")
P("Dataset: Pima Indians Diabetes Database (Kaggle: uciml/pima-indians-diabetes-database). Methods: K-Means and DBSCAN.",i=True)

H("1. Abstract")
P(f"This project looks for natural groups of patients using eight clinical attributes from the Pima Indians Diabetes Database ({S['n']} patients), without using the diabetes label during clustering. After cleaning impossible zero values and standardising the data, we compared K-Means and DBSCAN. K-Means with k = 3 gave three groups that are stable across random starts (mean adjusted Rand index {S['kmeans_stability_ari_mean']:.2f}) and clinically interpretable: a younger lower-risk group, an older group with many pregnancies, and a younger group with high BMI and insulin. The diabetes rate in these groups, which we checked only afterwards, was 13%, 52% and 52%, against 35% overall. The clustering structure is weak, however: the silhouette score is only {S['kmeans_silhouette']:.2f}, which means the groups overlap. DBSCAN found one dense group and {S['dbscan']['noise_points']} outlying patients ({100*S['dbscan']['noise_points']/S['n']:.1f}%), so in this dataset it works better as an outlier detector than as a segmentation tool.")

H("2. Introduction and objective")
P("Hospitals care for patients who differ widely in age, body measurements and blood tests. Grouping patients with similar clinical profiles can help plan screening, follow-up and awareness programmes. The objective is to discover patient groups from clinical attributes using unsupervised learning, and to judge honestly how well the groups are separated.")

H("3. Dataset")
P("The Pima Indians Diabetes Database comes from the National Institute of Diabetes and Digestive and Kidney Diseases. It has 768 records of women aged 21 or older of Pima Indian heritage, with eight attributes and an Outcome label (1 = diabetes, 0 = no diabetes). Outcome is NOT used for clustering. It is used only after clustering to check whether the groups mean something.")
B(["Pregnancies, Glucose, BloodPressure, SkinThickness, Insulin, BMI, DiabetesPedigreeFunction, Age.",
   f"Overall diabetes rate in the data: {100*S['overall_diabetes_rate']:.1f}%."])
P("Source note: we used a public copy of the same dataset (768 rows, same columns) because the Kaggle link needs a login. Please confirm against the Kaggle download before submission.",i=True)

H("4. Data preparation")
m=S["zeros_as_missing"]
P("Zero is not a possible value for glucose, blood pressure, skin thickness, insulin or BMI, so zeros in these columns were treated as missing values.")
B([f"{k}: {v} zero values ({100*v/S['n']:.1f}%)" for k,v in m.items()])
P("Missing values were filled with the column median. Insulin (about half) and skin thickness (about 30%) are heavily filled, so any pattern in those two columns should be read with care. All features were then standardised (mean 0, standard deviation 1) because K-Means and DBSCAN use distances and would otherwise be dominated by large-valued columns such as insulin.")
FIG("03_missing.png","Figure 1. Zero values treated as missing.",10)
FIG("01_distributions.png","Figure 2. Distribution of each feature after cleaning.",16)
FIG("02_correlation.png","Figure 3. Correlation between features.",11)

H("5. Methodology")
H("5.1 K-Means",2)
P("K-Means splits the patients into k groups by moving k centres until each patient is closest to its own centre. We tried k = 2 to 8 and compared inertia (elbow), silhouette score, Davies-Bouldin index and Calinski-Harabasz index. We used 10 random starts for each k.")
kt=km.round(3).rename(columns={"k":"k","inertia":"Inertia","silhouette":"Silhouette","davies_bouldin":"Davies-Bouldin","calinski_harabasz":"Calinski-Harabasz"})
TBL(kt)
FIG("04_kmeans_k_selection.png","Figure 4. Choosing k.",16)
P("How k = 3 was chosen: the silhouette score is almost the same for k = 2 and k = 3 (0.196 and 0.195) and then falls. The Davies-Bouldin index is better at k = 3 than k = 2. There is no sharp elbow. We chose k = 3 because it gives three balanced, interpretable groups. This is a judgement, not a result of the metrics alone, and k = 2 (a younger low-risk group and an older higher-risk group) is also reasonable.")
H("5.2 DBSCAN",2)
P("DBSCAN groups patients that sit in dense regions and marks isolated patients as noise. It needs two settings: eps (neighbourhood size) and min_samples. We used min_samples = 16 (about twice the number of features) and read eps from the k-distance plot, then searched a grid of eps from 1.0 to 3.0 and min_samples of 5, 8, 12 and 16.")
FIG("05_dbscan_kdistance.png","Figure 5. k-distance plot used to choose eps.",10)
P(f"Only two settings in the whole grid gave at least two clusters with under 25% noise, and in both the second cluster contained 4 to 5 patients (under 1%), which is a pocket of outliers and not a patient segment. We therefore used eps = {S['dbscan']['eps']}, min_samples = {S['dbscan']['min_samples']}, which gives one main cluster plus noise.")

H("6. Results")
H("6.1 K-Means segments",2)
pr=prof.rename(columns={"cluster":"Cluster","size":"Patients","DiabetesPedigreeFunction":"Pedigree","BloodPressure":"BP","SkinThickness":"Skin","diabetes_rate":"Diabetes rate"})
pr["Diabetes rate"]=(pr["Diabetes rate"]*100).round(0).astype(int).astype(str)+"%"
TBL(pr,fs=8)
FIG("07_kmeans_profile_heatmap.png","Figure 6. Average of each standardised feature in each cluster.",15)
B(["Cluster 0 (248 patients): older (average 46), many pregnancies (7.4), higher blood pressure. Diabetes rate 52%.",
   "Cluster 1 (335 patients): younger (average 26), fewer pregnancies, lower glucose and BMI (28.5). Diabetes rate 13%.",
   "Cluster 2 (185 patients): younger (average 29), highest BMI (39.2), skin thickness and insulin, high glucose. Diabetes rate 52%."])
FIG("08_kmeans_outcome_rate.png","Figure 7. Diabetes rate in each K-Means cluster. The label was not used to make the clusters.",9)
P(f"Stability: when K-Means was re-run with 10 different random starts, the clusters matched the original almost exactly (adjusted Rand index mean {S['kmeans_stability_ari_mean']:.3f}, minimum {S['kmeans_stability_ari_min']:.3f}).")
H("6.2 DBSCAN",2)
pd2=pdb.rename(columns={"cluster":"Cluster","size":"Patients","DiabetesPedigreeFunction":"Pedigree","BloodPressure":"BP","SkinThickness":"Skin","diabetes_rate":"Diabetes rate"})
pd2["Diabetes rate"]=(pd2["Diabetes rate"]*100).round(0).astype(int).astype(str)+"%"
TBL(pd2,fs=8)
P("Cluster -1 is the noise group. The 88 noise patients have higher glucose (146), much higher insulin (266) and a higher pedigree function (0.81) than the main group, and their diabetes rate is 55% against 32% in the main group. DBSCAN is therefore picking out unusual high-risk patients.")
H("6.3 Visual comparison",2)
FIG("06_pca_clusters.png",f"Figure 8. The same patients on the first two principal components (together {100*S['pca_variance_2']:.0f}% of the variance).",16)

H("7. Discussion")
B(["K-Means gives more useful segments here. The three groups are different in age, body measurements and diabetes rate, and they are very stable.",
   f"But the groups overlap. The silhouette score of {S['kmeans_silhouette']:.2f} is low, and in the plot the three clusters touch each other. The data is better described as a continuous spread of patients that K-Means cuts into regions, not as three clearly separate types.",
   "DBSCAN does not find several groups because there are no clearly separated dense regions. It is better used to flag outliers.",
   "The diabetes rates were not used to build the clusters, so the fact that they differ (13% against 52%) is an independent sign that the groups carry clinical meaning."])
H("8. Limitations")
B(["The data covers only women aged 21 or older of Pima heritage, so results cannot be assumed to hold for other patients.",
   "768 patients is a small sample.",
   "Insulin and skin thickness are heavily imputed with the median, which may blur real differences.",
   "Clusters describe similarity of measurements. They are not diagnoses and must not be used for medical decisions."])
H("9. Conclusion")
P("Using K-Means with k = 3, the patients fall into three stable and interpretable groups with very different diabetes rates (13%, 52%, 52%), although the separation between groups is weak. DBSCAN found one main group and 88 unusual patients. For patient segmentation on this dataset, K-Means is the better method, and DBSCAN is useful for spotting outliers. Future work could try Gaussian mixture models, hierarchical clustering, or better handling of missing insulin values.")
H("10. How to run")
P("pip install -r requirements.txt, then python3 patient_segmentation.py. Figures are saved in figures/ and tables in results/.")
H("11. References")
for r in ["Smith, J.W., Everhart, J.E., Dickson, W.C., Knowler, W.C., Johannes, R.S. (1988). Using the ADAP learning algorithm to forecast the onset of diabetes mellitus. Proc. Symposium on Computer Applications in Medical Care.",
 "Pedregosa, F. et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12.",
 "Ester, M., Kriegel, H.-P., Sander, J., Xu, X. (1996). A density-based algorithm for discovering clusters in large spatial databases with noise (DBSCAN). KDD.",
 "Rousseeuw, P.J. (1987). Silhouettes: a graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, 20.",
 "Kaggle: https://www.kaggle.com/datasets/uciml/pima-indians-diabetes-database"]:
    d.add_paragraph(r,style="List Bullet")
d.save("Patient_Segmentation_Report.docx")
