"""Builds the Secure and Resilient Virtualized Data Centre Design report (DOCX + PDF). Two passes so that the table of contents carries real page numbers."""
import re, subprocess, pathlib, sys
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docbuilder import Report, _set_font
import content as C

HERE = pathlib.Path(__file__).parent
TITLE = "SECURE AND RESILIENT VIRTUALIZED DATA CENTRE DESIGN"
TITLE_TC = "Secure and Resilient Virtualized Data Centre Design"
CHAPTERS = ["Introduction", "Objectives", "Tools and Technologies", "Architecture Diagram", "Methodology", "Results", "Conclusion", "Reference"]
STUDENTS = [("111624102241", "A.VISHNU VARDHAN"), ("111624102242", "VISHWAJITH R"), ("111624102247", "YUVARAJ U")]
GUIDE = "Dr. J GNANA JESLIN"

from abstract import ABSTRACT


def centre(r, text, size=12, bold=True, italic=False, before=0, after=0, spacing=1.15):
    return r.para(text, align="c", size=size, bold=bold, italic=italic, before=before, after=after, spacing=spacing)


def cover(r):
    centre(r, TITLE, 17, before=40, after=22, spacing=1.3)
    centre(r, "A PROJECT REPORT", 13, after=26)
    centre(r, "Submitted by", 12, italic=True, after=22)
    t = r.doc.add_table(rows=1 + len(STUDENTS), cols=2)
    Report._fix_widths(t, (6.8, 7.4))
    for cell, txt, sz in zip(t.rows[0].cells, ("REGNO", "NAME"), (14, 14)):
        cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6); _set_font(p.add_run(txt), 14, True)
    for row, (reg, name) in zip(t.rows[1:], STUDENTS):
        for cell, txt in zip(row.cells, (reg, name)):
            cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3); _set_font(p.add_run(txt), 12, True)
    centre(r, "", after=14)
    centre(r, "in partial fulfillment for the award of the degree", 12, italic=True, after=6)
    centre(r, "of", 12, italic=True, after=6)
    centre(r, "BACHELOR OF ENGINEERING", 12, after=6)
    centre(r, "in", 12, italic=True, after=6)
    centre(r, "COMPUTER SCIENCE AND ENGINEERING", 12, after=26)
    centre(r, "R. M. K. COLLEGE OF ENGINEERING AND TECHNOLOGY", 12, after=2)
    centre(r, "(An Autonomous Institution)", 10, after=14)
    centre(r, "PUDUVOYAL", 12, after=10)
    p = r.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(HERE / "assets/crest_cover.png"), width=Cm(2.6))
    centre(r, "OCTOBER 2026", 12, before=8)
    r.page_break()


def bonafide(r):
    t = r.doc.add_table(rows=1, cols=3)
    Report._fix_widths(t, (2.4, 10.7, 2.4))
    c0, c1, c2 = t.rows[0].cells
    c0.paragraphs[0].add_run().add_picture(str(HERE / "assets/crest_cover.png"), width=Cm(2.2))
    c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    c2.paragraphs[0].add_run().add_picture(str(HERE / "assets/iso.png"), width=Cm(2.1))
    def line(txt, size, color, bold=True, first=False):
        p = c1.paragraphs[0] if first else c1.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
        _set_font(p.add_run(txt), size, bold, color=color)
    line("R.M.K. COLLEGE OF ENGINEERING AND TECHNOLOGY", 13.5, "C00000", False, True)
    line("(An Autonomous Institution)", 11, "C55A11", False)
    for txt in ("R.S.M. Nagar, PUDUVOYAL-601 206", "Approved by AICTE, New Delhi /Affiliated to Anna University, Chennai",
                "Accredited by NBA, New Delhi (All Eligible Courses)/ NAAC with ‘A’ Grade", "An ISO 21001:2018 Certified Institution"):
        line(txt, 7.5, "843C0C")
    centre(r, "", after=18)
    centre(r, "BONAFIDE CERTIFICATE", 14, after=18)
    names = [n for _, n in STUDENTS]; regs = [g for g, _ in STUDENTS]
    nt = ", ".join(names[:-1]) + " and " + names[-1]; rt = ", ".join(regs[:-1]) + " and " + regs[-1]
    r.para(f"This is to certify that **{nt}**, bearing **Register Numbers {rt}** respectively, of **3rd Year, 5th Semester, "
           f"Department of Computer Science and Engineering**, have successfully completed the mini project entitled "
           f"**“{TITLE_TC}”** during the academic year **2026–2027** under my guidance and supervision.", spacing=1.5)
    r.para("The work submitted in this report is a bonafide record of the work carried out by the student/team as part of the "
           "academic requirements of the course.", spacing=1.5)
    centre(r, "", after=40)
    t = r.doc.add_table(rows=1, cols=2)
    Report._fix_widths(t, (8.2, 6.8))
    for cell, txt in zip(t.rows[0].cells, ("Name of the Student", "Project Guide")):
        cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_font(p.add_run(txt), 12, True)
    for k, n in enumerate(names):
        row = t.add_row().cells
        row[0].text = ""; p = row[0].paragraphs[0]; p.paragraph_format.space_after = Pt(2); _set_font(p.add_run(n), 12)
        if k == 0:
            row[1].text = ""; q = row[1].paragraphs[0]; q.paragraph_format.space_after = Pt(2); _set_font(q.add_run(GUIDE), 12)
    r.page_break()


def abstract(r):
    centre(r, "ABSTRACT", 14, after=14)
    r.para(ABSTRACT, spacing=1.5)
    r.page_break()


def toc(r, pages):
    centre(r, "TABLE OF CONTENTS", 14, after=14)
    t = r.doc.add_table(rows=1, cols=3); t.style = "Table Grid"; t.autofit = False
    w = (Cm(1.6), Cm(9.0), Cm(4.4))
    for i, (c, txt) in enumerate(zip(t.rows[0].cells, ("S.No", "Section", "Page No."))):
        c.width = w[i]; c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(txt), 12, True)
    for i, name in enumerate(CHAPTERS, 1):
        cells = t.add_row().cells
        for j, (c, txt) in enumerate(zip(cells, (str(i), name, str(pages["chapters"].get(i, ""))))):
            c.width = w[j]; c.text = ""; p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
            _set_font(p.add_run(txt), 12)
    Report._fix_widths(t, (1.6, 9.0, 4.4))
    r.para("", after=6); r.page_break()


def build(out, pages):
    r = Report()
    cover(r); bonafide(r); abstract(r); toc(r, pages)
    r.start_body()
    C.chapter1(r); C.chapter2(r); C.chapter3(r); C.chapter4(r)
    C.chapter5(r); C.chapter6(r); C.chapter7(r); C.chapter8(r)
    r.doc.core_properties.title = "Secure and Resilient Virtualized Data Centre Design - Project Report"
    r.doc.core_properties.author = "A.Vishnu Vardhan; Vishwajith R; Yuvaraj U"
    r.save(out)


def to_pdf(docx, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)


def locate(pdf):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
    texts = [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdf), "-"], capture_output=True, text=True).stdout for i in range(1, n + 1)]
    start = next(i for i, t in enumerate(texts) if re.search(r"^\s*1\. INTRODUCTION\s*$", t, re.M))
    heads = {1: "1. INTRODUCTION", 2: "2. OBJECTIVES", 3: "3. TOOLS AND TECHNOLOGIES", 4: "4. ARCHITECTURE DIAGRAM", 5: "5. METHODOLOGY", 6: "6. RESULTS", 7: "7. CONCLUSION", 8: "8. REFERENCE"}
    pages = {"chapters": {}, "total": n, "body_start": start + 1}
    for k, h in heads.items():
        for i in range(start, n):
            if re.search(rf"^\s*{re.escape(h)}\s*$", texts[i], re.M):
                pages["chapters"][k] = i - start + 1; break
    return pages


if __name__ == "__main__":
    out = HERE / "output"; out.mkdir(exist_ok=True)
    docx = out / "Secure_Resilient_Virtualized_Data_Centre_Design_Report.docx"; pdf = out / "Secure_Resilient_Virtualized_Data_Centre_Design_Report.pdf"
    build(docx, {"chapters": {i: 99 for i in range(1, 9)}}); to_pdf(docx, out)
    pages = locate(pdf); build(docx, pages); to_pdf(docx, out)
    p2 = locate(pdf)
    if p2["chapters"] != pages["chapters"]:
        build(docx, p2); to_pdf(docx, out); p2 = locate(pdf)
    print("total pdf pages", p2["total"], "body starts at", p2["body_start"], p2["chapters"])
