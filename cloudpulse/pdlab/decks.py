"""Three review decks (12 slides each) for the CloudPulse PD Lab III project.
Design: warm paper background, Cambria titles, Calibri body, one terracotta accent, left-aligned layouts,
real tables / diagrams / screenshots. No icons, gradients or stock graphics."""
import pathlib
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from lxml import etree

HERE = pathlib.Path(__file__).parent
A = str(HERE.parent / "internship" / "assets") + "/"
OUT = HERE / "output"; OUT.mkdir(exist_ok=True)

PAPER = RGBColor(0xFA, 0xF7, 0xF2)
INK = RGBColor(0x1E, 0x2A, 0x32)
MUTED = RGBColor(0x6B, 0x72, 0x78)
RULE = RGBColor(0xD9, 0xD2, 0xC6)
ACCENT = RGBColor(0xB5, 0x53, 0x2B)     # terracotta
TEAL = RGBColor(0x2F, 0x5D, 0x62)
SOFT = RGBColor(0xF1, 0xEA, 0xDF)       # slightly darker paper for panels
PH = RGBColor(0xB0, 0x2A, 0x7A)         # placeholder colour
HEAD, BODY = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)
LM = Inches(0.85)                         # left margin
CW = Inches(11.6)                         # content width

TEAM = [("Tamil Mani B", "111624102217"), ("Theja J G", "111624102225"),
        ("Varunesh D", "111624102235"), ("Viswajith R", "111624102242")]
PROJECT = "CloudPulse"
SUB = "A cloud cost monitoring and optimisation platform that proves its savings"


# ------------------------------------------------------------------ primitives
def new_deck():
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    return prs


def blank(prs, review, n, total=12, section=None, title=None, footer=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER
    if title:
        if section:
            text(s, LM, Inches(0.5), Inches(8), Inches(0.3), section.upper(), 11, ACCENT, True, font=BODY, spacing=2)
        text(s, LM, Inches(0.78), CW, Inches(0.8), title, 30, INK, True, font=HEAD)
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, LM, Inches(1.5), Inches(0.7), Emu(28000))
        ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background(); ln.shadow.inherit = False
    if footer:
        text(s, LM, Inches(7.02), Inches(8), Inches(0.3),
             f"{PROJECT}  ·  Product Development Lab III  ·  Review {review}", 10, MUTED, font=BODY)
        text(s, Inches(11.4), Inches(7.02), Inches(1.1), Inches(0.3), f"{n} / {total}", 10, MUTED, font=BODY, align=PP_ALIGN.RIGHT)
        hr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, LM, Inches(6.97), CW, Emu(9000))
        hr.fill.solid(); hr.fill.fore_color.rgb = RULE; hr.line.fill.background(); hr.shadow.inherit = False
    return s


def text(s, x, y, w, h, t, size=18, color=INK, bold=False, italic=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=None):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line: p.line_spacing = line
        r = p.add_run(); r.text = ln
        f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return tb


def rich(s, x, y, w, h, paras, size=18, gap=8, color=INK, line=1.08):
    """paras: list of items; item = str or (lead, rest). Lead is set in bold accent colour."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = line
        lead, rest = (it if isinstance(it, tuple) else (None, it))
        if lead:
            r = p.add_run(); r.text = lead + " "; f = r.font; f.name = BODY; f.size = Pt(size); f.bold = True; f.color.rgb = ACCENT
        r = p.add_run(); r.text = rest; f = r.font; f.name = BODY; f.size = Pt(size); f.color.rgb = color
    return tb


def dash_list(s, x, y, w, h, items, size=17, gap=9):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.08
        r = p.add_run(); r.text = "–  "; f = r.font; f.name = BODY; f.size = Pt(size); f.color.rgb = ACCENT; f.bold = True
        lead, rest = (it if isinstance(it, tuple) else (None, it))
        if lead:
            r = p.add_run(); r.text = lead + " "; f = r.font; f.name = BODY; f.size = Pt(size); f.bold = True; f.color.rgb = INK
        r = p.add_run(); r.text = rest; f = r.font; f.name = BODY; f.size = Pt(size); f.color.rgb = INK
    return tb


def panel(s, x, y, w, h, fill=SOFT):
    p = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    p.fill.solid(); p.fill.fore_color.rgb = fill; p.line.fill.background(); p.shadow.inherit = False
    return p


def rule(s, x, y, w, color=RULE, weight=9000):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Emu(weight))
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); r.shadow.inherit = False; return r


def table(s, x, y, col_w, header, rows, size=13, row_h=0.42, head_color=MUTED, bold_first=False, bold_last=False):
    shape = s.shapes.add_table(len(rows) + 1, len(header), x, y, Inches(sum(col_w)), Inches(row_h * (len(rows) + 1)))
    t = shape.table
    tblPr = t._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(col_w):
        t.columns[i].width = Inches(w)

    def cell(r, c, txt, hdr=False, bold=False):
        ce = t.cell(r, c); ce.fill.background()
        ce.margin_left = Inches(0.06); ce.margin_right = Inches(0.08); ce.margin_top = Inches(0.04); ce.margin_bottom = Inches(0.04)
        ce.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = ce.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; run = p.add_run(); run.text = str(txt)
        f = run.font; f.name = BODY; f.size = Pt(size - 1.5 if hdr else size); f.bold = hdr or bold
        f.color.rgb = head_color if hdr else INK
        if hdr:
            run._r.get_or_add_rPr().set("spc", "80")
        tcPr = ce._tc.get_or_add_tcPr()
        for side in ("a:lnL", "a:lnR", "a:lnT"):
            ln = etree.SubElement(tcPr, qn(side)); ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
        ln = etree.SubElement(tcPr, qn("a:lnB")); ln.set("w", "9525" if not hdr else "15875")
        sf = etree.SubElement(ln, qn("a:solidFill")); clr = etree.SubElement(sf, qn("a:srgbClr")); clr.set("val", "1E2A32" if hdr else "D9D2C6")
        # schema order: borders must precede the fill
        fill = tcPr.find(qn("a:noFill"))
        if fill is not None:
            tcPr.remove(fill); tcPr.append(fill)

    for c, h in enumerate(header):
        cell(0, c, h.upper(), hdr=True)
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row):
            cell(r, c, v, bold=(bold_first and c == 0) or (bold_last and r == len(rows)))
    for r in range(len(rows) + 1):
        t.rows[r].height = Inches(row_h)
    return shape


def pic(s, path, x, y, w=None, h=None, border=True):
    kw = {"width": w} if w is not None else {"height": h}
    p = s.shapes.add_picture(path, x, y, **kw)
    if border:
        p.line.color.rgb = RULE; p.line.width = Pt(0.75)
    return p


def caption(s, x, y, w, t):
    return text(s, x, y, w, Inches(0.3), t, 11, MUTED, italic=True)


def note(s, x, y, w, t, size=12):
    """small hand-written-feeling remark in the accent colour"""
    return text(s, x, y, w, Inches(0.5), t, size, ACCENT, italic=True, font=HEAD)


# ------------------------------------------------------------------ shared slides
def title_slide(prs, review, subtitle_line):
    s = blank(prs, review, 1, footer=False)
    panel(s, Inches(8.45), 0, Inches(4.883), H, SOFT)
    s.shapes.add_picture(A + "crest_slide.png", LM, Inches(0.55), height=Inches(0.8))
    text(s, Inches(1.85), Inches(0.62), Inches(6.2), Inches(0.7),
         ["R.M.K. College of Engineering and Technology", "Department of Computer Science and Engineering"], 12, MUTED, line=1.15)
    text(s, LM, Inches(2.05), Inches(7.4), Inches(0.4), subtitle_line.upper(), 12, ACCENT, True, spacing=3)
    text(s, LM, Inches(2.5), Inches(7.4), Inches(1.2), PROJECT, 66, INK, True, font=HEAD)
    text(s, LM, Inches(3.75), Inches(7.0), Inches(1.2), SUB, 22, INK, font=HEAD, italic=True, line=1.1)
    rule(s, LM, Inches(5.25), Inches(0.7), ACCENT, 28000)
    text(s, LM, Inches(5.5), Inches(7.2), Inches(0.9),
         ["Product Development Lab III  ·  24AI511", "Review date: [DD / 10 / 2026]     Supervisor: [name]"], 14, MUTED, line=1.3)
    # right column: team
    text(s, Inches(8.95), Inches(1.6), Inches(3.9), Inches(0.35), "TEAM", 11, ACCENT, True, spacing=3)
    y = Inches(2.1)
    for n, r in TEAM:
        text(s, Inches(8.95), y, Inches(3.9), Inches(0.4), n, 20, INK, True, font=HEAD)
        text(s, Inches(8.95), y + Inches(0.4), Inches(3.9), Inches(0.3), r, 13, MUTED)
        y += Inches(0.95)
    text(s, Inches(8.95), Inches(6.55), Inches(3.9), Inches(0.3), "October 2026", 12, MUTED)
    return s


def agenda(prs, review, items, n=2):
    s = blank(prs, review, n, title="What we will walk through", section="Agenda")
    y = Inches(2.0)
    colx = [LM, LM + Inches(5.95)]
    per = (len(items) + 1) // 2
    for i, (name, slide_no) in enumerate(items):
        col, row = divmod(i, per)
        x = colx[col]; yy = y + Inches(0.8) * row
        text(s, x, yy, Inches(0.7), Inches(0.5), f"{i + 1:02d}", 22, ACCENT, True, font=HEAD)
        text(s, x + Inches(0.75), yy + Inches(0.04), Inches(4.4), Inches(0.5), name, 20, INK)
        text(s, x + Inches(4.95), yy + Inches(0.08), Inches(0.7), Inches(0.5), f"p. {slide_no}", 12, MUTED, align=PP_ALIGN.RIGHT)
        rule(s, x, yy + Inches(0.62), Inches(5.65))
    return s


ABSTRACT = ("Organisations and students who run workloads on Azure and AWS keep paying for resources nobody uses: virtual machines shut "
            "down the wrong way that still bill for compute, disks left behind after a project, public IPs attached to nothing. "
            "Cost consoles show how much was spent, but not who owns the waste, what to fix first, or whether a fix really saved money.")
ABSTRACT2 = ("CloudPulse connects to an Azure subscription with a read-only identity, finds the waste and spend anomalies, forecasts budget "
             "breaches, assigns every finding to the person who created the resource, and counts a saving only after seven days of billed "
             "cost confirm it. In the demo workspace it reports $262.94 a month still to fix and $33.44 a month already verified.")


def abstract_slide(prs, review, n):
    s = blank(prs, review, n, title="Abstract", section="Summary")
    text(s, LM, Inches(2.0), Inches(6.6), Inches(3.4), ABSTRACT, 19, INK, line=1.2)
    text(s, LM, Inches(4.45), Inches(6.6), Inches(2.4), ABSTRACT2, 19, INK, line=1.2)
    panel(s, Inches(8.1), Inches(2.0), Inches(4.35), Inches(4.5))
    text(s, Inches(8.45), Inches(2.25), Inches(3.7), Inches(0.3), "IN NUMBERS  (DEMO WORKSPACE)", 11, ACCENT, True, spacing=2)
    y = Inches(2.75)
    for big, small in (("$262.94", "a month of open waste, 6 findings"), ("$33.44", "a month verified on the bill"),
                       ("94 %", "estimate accuracy of verified savings"), ("27", "automated tests passing")):
        text(s, Inches(8.45), y, Inches(3.7), Inches(0.6), big, 30, INK, True, font=HEAD)
        text(s, Inches(8.45), y + Inches(0.55), Inches(3.7), Inches(0.3), small, 13, MUTED)
        y += Inches(0.92)
    return s


OBJECTIVES = [
    ("Collect", "cost, inventory, usage, advice and audit data from Azure with a least-privilege, read-only identity."),
    ("Detect", "unattached disks, unused public IPs, orphaned snapshots, stopped-but-billed, idle and oversized VMs."),
    ("Explain", "abnormal daily spend and forecast month-end spend and the day a budget will be breached."),
    ("Assign", "every finding to the person who created the resource, using the Activity Log and tags."),
    ("Verify", "each saving against the next bill, and withdraw it if the waste comes back."),
    ("Prove", "the result with an Azure test environment that plants real waste and 27 automated tests."),
]


def objectives_slide(prs, review, n, with_status=None):
    s = blank(prs, review, n, title="Objectives", section="What we set out to do")
    y = Inches(1.95)
    for i, (lead, rest) in enumerate(OBJECTIVES):
        yy = y + Inches(0.78) * i
        text(s, LM, yy, Inches(0.7), Inches(0.6), f"{i + 1}", 30, ACCENT, True, font=HEAD)
        rich(s, LM + Inches(0.7), yy + Inches(0.04), Inches(10.7), Inches(0.7), [(lead, rest)], 18, 0)
        rule(s, LM + Inches(0.7), yy + Inches(0.66), Inches(10.9))
    return s


# ====================================================================== REVIEW 1
def review1():
    prs = new_deck(); R = 1
    title_slide(prs, R, "Review I")
    agenda(prs, R, [("The pain point", 3), ("Abstract", 4), ("Objectives", 5), ("Literature review", "6 – 8"),
                    ("Proposed solution", 9), ("Functional requirements", 10), ("Non-functional requirements and tools", 11),
                    ("References", 12)])

    # 3 pain point
    s = blank(prs, R, 3, title="The pain point: waste that nobody owns", section="Why this project")
    text(s, LM, Inches(1.95), Inches(5.3), Inches(1.5),
         "A forgotten resource keeps billing quietly. Each one is small; together they decide whether a student's free credit or a "
         "team's monthly budget survives.", 18, INK, line=1.2)
    rows = [["VM shut down from inside the OS", "Looks stopped, but compute is still billed", "$70.08 / month"],
            ["Disk left after a project ended", "Billed per gigabyte until it is deleted", "$19.71 / month"],
            ["Static public IP on nothing", "Billed hourly even when unattached", "$3.65 / month"],
            ["Snapshot of a deleted disk", "Nobody remembers why it exists", "$3.20 / month"]]
    table(s, Inches(6.55), Inches(1.95), [2.55, 2.2, 1.3], ["Leftover", "Why it hurts", "Example cost"], rows, size=12, row_h=0.62)
    rich(s, LM, Inches(3.6), Inches(5.3), Inches(2.8), [
        ("No owner.", "Tags are often missing, so nobody knows who can switch it off."),
        ("No priority.", "A $200 VM and a $3 IP look equally urgent in a long list."),
        ("No proof.", "A closed ticket is not evidence. Only a lower bill is.")], 17, 10)
    note(s, Inches(6.55), Inches(5.35), Inches(5.7), "Example costs are taken from our demo workspace, not from a customer.", 12)

    abstract_slide(prs, R, 4)
    objectives_slide(prs, R, 5)

    # 6 literature review 1
    s = blank(prs, R, 6, title="Literature review: what exists today", section="Literature review 1 / 3")
    cols = [("Native consoles", "Azure Cost Management, Azure Advisor, AWS Cost Explorer",
             "Accurate and free. One provider each, and they stop at a report or a generic tip."),
            ("Commercial platforms", "Multi-cloud dashboards, cost allocation, budgets",
             "Broad, but closed: the scoring is not visible and broad access is usually required."),
            ("Home-made scripts", "List idle resources with the CLI or an SDK",
             "Cheap and flexible, but one-off: no workflow, history or verification.")]
    x = LM
    for t, ex, v in cols:
        text(s, x, Inches(2.0), Inches(3.6), Inches(0.4), t, 21, INK, True, font=HEAD)
        rule(s, x, Inches(2.5), Inches(3.5), ACCENT, 19000)
        text(s, x, Inches(2.7), Inches(3.5), Inches(0.9), ex, 14, MUTED, italic=True, line=1.15)
        text(s, x, Inches(3.55), Inches(3.5), Inches(2.0), v, 17, INK, line=1.2)
        x += Inches(3.95)
    panel(s, LM, Inches(5.45), CW, Inches(1.2))
    rich(s, LM + Inches(0.3), Inches(5.65), Inches(11.0), Inches(0.9), [
        ("FinOps practice.", "The FinOps Foundation frames cloud cost as shared accountability: inform, optimise, operate. Visibility alone is not the goal; "
                             "a problem is solved when its owner acts and the saving shows on the bill [3], [9].")], 16, 0)

    # 7 literature review 2
    s = blank(prs, R, 7, title="Literature review: techniques we build on", section="Literature review 2 / 3")
    rows = [["Cloud model and billing", "Armbrust et al. 2010 [1]; NIST SP 800-145 [2]", "Pay-as-you-go and elasticity explain why waste appears"],
            ["Anomaly detection", "Chandola, Banerjee, Kumar 2009 [5]; Iglewicz and Hoaglin 1993 [6]", "Median / MAD z-score, cut-off 3.5, same-weekday baseline"],
            ["Forecasting", "Holt 1957 / 2004 [7]; Hyndman and Athanasopoulos [8]", "Holt linear smoothing with weekday factors and an 80 % range"],
            ["FinOps practice", "FinOps Foundation [3]; Storment and Fuller [9]", "Inform, optimise, operate as the product loop"],
            ["Identity and access", "OAuth 2.0, RFC 6749 [10]; Azure RBAC [11]", "Service principal, client-credentials token, Reader roles"],
            ["Cost-data format", "FinOps FOCUS specification [12]", "Import of multi-cloud bills in one open format"]]
    table(s, LM, Inches(2.0), [2.6, 4.5, 4.5], ["Area", "Source", "What we take from it"], rows, size=13, row_h=0.62, bold_first=True)

    # 8 gap table
    s = blank(prs, R, 8, title="Literature review: the gap we address", section="Literature review 3 / 3")
    rows = [["Spend dashboards and breakdown", "Yes", "Yes", "Rarely", "Yes"],
            ["Idle and oversized resource detection", "Partial", "Yes", "Partial", "Yes"],
            ["Creator of a resource from the audit log", "Manual", "Varies", "No", "Automatic"],
            ["Task workflow: assign, start, fix, dismiss", "No", "Varies", "No", "Yes"],
            ["One explainable priority score", "No", "Varies", "No", "Yes"],
            ["Saving checked against the next bill", "No", "Rarely", "No", "Yes"],
            ["Saving withdrawn when waste returns", "No", "Rarely", "No", "Yes"],
            ["Read-only access, open algorithms", "n/a", "Varies", "Depends", "Yes"]]
    table(s, LM, Inches(1.95), [4.6, 1.7, 1.9, 1.5, 1.9], ["Capability", "Native consoles", "Commercial tools", "Scripts", "CloudPulse"],
          rows, size=13, row_h=0.47)
    note(s, LM, Inches(6.35), Inches(11), "Gap: owner attribution, prioritisation and verified savings are rarely combined in one transparent, read-only tool.", 14)

    # 9 proposed solution loop
    s = blank(prs, R, 9, title="Proposed solution: the cost-to-action loop", section="Our approach")
    steps = [("Signal", "Read the bill, inventory, metrics, advice and audit log"),
             ("Context", "Price each signal per month from billed cost"),
             ("Finding", "One de-duplicated, scored task with evidence"),
             ("Ownership", "Creator from the Activity Log, tags second"),
             ("Action", "Assign, start, fix, dismiss or snooze"),
             ("Verification", "Compare 7 days of bill before and after")]
    bw, gap = Inches(1.78), Inches(0.18)
    for i, (t, d) in enumerate(steps):
        x = LM + (bw + gap) * i
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.2), bw, Inches(2.2))
        b.fill.solid(); b.fill.fore_color.rgb = SOFT if i < 5 else ACCENT; b.line.fill.background(); b.shadow.inherit = False
        text(s, x + Inches(0.15), Inches(2.35), bw - Inches(0.3), Inches(0.4), f"{i + 1}", 14, ACCENT if i < 5 else PAPER, True)
        text(s, x + Inches(0.15), Inches(2.75), bw - Inches(0.3), Inches(0.5), t, 20, INK if i < 5 else PAPER, True, font=HEAD)
        text(s, x + Inches(0.15), Inches(3.3), bw - Inches(0.3), Inches(1.1), d, 13, INK if i < 5 else PAPER, line=1.15)
    rich(s, LM, Inches(4.9), Inches(7.2), Inches(1.9), [
        ("Read-only.", "CloudPulse never creates, changes or deletes anything in a cloud account. It suggests the command; a person runs it."),
        ("Honest savings.", "A saving counts only after the bill drops, and the finding reopens if the resource comes back.")], 16, 9)
    note(s, Inches(8.4), Inches(5.0), Inches(4.0), "Everything on the next slides exists to serve this loop.", 14)

    # 10 functional requirements
    s = blank(prs, R, 10, title="Functional requirements", section="Requirements 1 / 2")
    rows = [["FR1", "Sign in with e-mail and password; read-only demo workspace", "FR7", "Anomalies with drivers; month forecast and budget breach date"],
            ["FR2", "Create workspaces; connect Azure (service principal, CLI, managed identity)", "FR8", "Finding workflow: assign, start, resolve, dismiss, snooze, comment"],
            ["FR3", "Read daily cost by service and by resource", "FR9", "7-day before / after verification and a savings ledger"],
            ["FR4", "Read inventory, VM power state and VM metrics", "FR10", "Import FOCUS bills; set a monthly budget"],
            ["FR5", "Detect waste and rank it by priority", "FR11", "Notifications, audit log and CSV export"],
            ["FR6", "Find the creator of each resource", "", ""]]
    table(s, LM, Inches(2.0), [0.7, 4.9, 0.8, 5.2], ["ID", "Requirement", "ID", "Requirement"], rows, size=13, row_h=0.7, bold_first=True)

    # 11 NFR + tools
    s = blank(prs, R, 11, title="Non-functional requirements and planned tools", section="Requirements 2 / 2")
    text(s, LM, Inches(1.95), Inches(5.6), Inches(0.4), "Qualities we will hold the system to", 18, INK, True, font=HEAD)
    dash_list(s, LM, Inches(2.5), Inches(5.6), Inches(4.2), [
        ("Security:", "least privilege, encrypted secrets, hashed passwords, rate limit"),
        ("Reliability:", "retry on 429 / 5xx; one failing source must not stop a sync"),
        ("Explainability:", "every score can be recomputed by hand"),
        ("Portability:", "laptop, Docker, Render, Azure App Service"),
        ("Usability:", "no framework to install; works from 400 px phones"),
        ("Cost:", "free-tier hosting; test workload about $2.3 a day")], 15, 8)
    text(s, Inches(7.0), Inches(1.95), Inches(5.4), Inches(0.4), "Tools we plan to use", 18, INK, True, font=HEAD)
    rows = [["Language and API", "Python 3, FastAPI"], ["Data", "SQLAlchemy, SQLite / PostgreSQL"],
            ["Interface", "HTML, CSS, JavaScript, SVG charts"], ["Cloud access", "Azure CLI, ARM REST APIs, azure-identity"],
            ["Packaging", "Docker, GitHub Actions"], ["Testing", "pytest"]]
    table(s, Inches(7.0), Inches(2.5), [1.8, 3.6], ["Layer", "Choice"], rows, size=13, row_h=0.55, bold_first=True)

    # 12 references
    s = blank(prs, R, 12, title="References and what comes next", section="Closing")
    refs = ["[1] M. Armbrust et al., “A view of cloud computing,” Commun. ACM, 53(4), 2010.",
            "[2] P. Mell and T. Grance, NIST SP 800-145, 2011.",
            "[3] FinOps Foundation, “What is FinOps?”  finops.org",
            "[5] V. Chandola, A. Banerjee, V. Kumar, “Anomaly detection: a survey,” ACM Comput. Surv., 41(3), 2009.",
            "[6] B. Iglewicz and D. Hoaglin, How to Detect and Handle Outliers, 1993.",
            "[7] C. C. Holt, Int. J. Forecasting, 20(1), 2004.",
            "[8] R. Hyndman and G. Athanasopoulos, Forecasting: Principles and Practice, 3rd ed.",
            "[9] J. Storment and M. Fuller, Cloud FinOps, 2nd ed., O'Reilly, 2023.",
            "[10] RFC 6749, The OAuth 2.0 Authorization Framework.   [11] Microsoft Learn, Azure built-in roles.",
            "[12] FinOps Foundation, FOCUS specification.  focus.finops.org"]
    text(s, LM, Inches(2.0), Inches(7.4), Inches(4.8), refs, 13, INK, line=1.3)
    panel(s, Inches(8.7), Inches(2.0), Inches(3.75), Inches(4.5))
    text(s, Inches(9.0), Inches(2.25), Inches(3.2), Inches(0.3), "NEXT: REVIEW II", 11, ACCENT, True, spacing=3)
    dash_list(s, Inches(9.0), Inches(2.75), Inches(3.2), Inches(3.5), [
        "System architecture", "Azure credential flow", "Module-by-module description"], 16, 10)
    text(s, Inches(9.0), Inches(5.3), Inches(3.2), Inches(0.9), "Thank you. Questions are welcome.", 16, INK, italic=True, font=HEAD)
    prs.save(OUT / "PD_Lab_III_Review_1.pptx")


# ====================================================================== REVIEW 2
def architecture_slide(prs, R, n):
    s = blank(prs, R, n, title="System architecture", section="Design")
    pic(s, A + "fig_architecture.png", Inches(2.9), Inches(1.7), h=Inches(5.15), border=False)
    rich(s, LM, Inches(2.0), Inches(1.95), Inches(4.8), [
        ("Client.", "Single-page browser app."),
        ("Server.", "One Python process: API, scheduler, analytics."),
        ("Data.", "SQL database, encrypted credential store."),
        ("Azure.", "Only read, with a Bearer token.")], 13, 9)
    caption(s, Inches(2.9), Inches(6.72), Inches(9), "Numbered flows: sign-in, credential storage, scheduled sync, token request, token response, ARM reads, results stored.")


def review2():
    prs = new_deck(); R = 2
    title_slide(prs, R, "Review II")
    agenda(prs, R, [("Abstract", 3), ("Objectives", 4), ("Literature review", 5), ("Pain points and design answers", 6),
                    ("System architecture", 7), ("Azure credential architecture", 8), ("Module description", "9 – 11"),
                    ("What the system reports", 12)])
    abstract_slide(prs, R, 3)
    objectives_slide(prs, R, 4)

    # 5 literature
    s = blank(prs, R, 5, title="Literature review in one page", section="Background")
    for i, (h, body) in enumerate([
            ("What exists", "Native consoles report cost; commercial platforms add allocation behind closed scoring; scripts give one-off lists [3], [4]."),
            ("What we borrow", "Robust z-score with a weekday baseline [5], [6]; Holt smoothing for forecasts [7], [8]; the FinOps loop [3], [9]; OAuth 2.0 and RBAC for access [10], [11]."),
            ("What is missing", "Owner attribution, one explainable priority and savings verified on the bill, in a read-only tool.")]):
        x = LM + Inches(3.95) * i
        text(s, x, Inches(2.0), Inches(3.6), Inches(0.5), h, 22, ACCENT if i == 2 else INK, True, font=HEAD)
        rule(s, x, Inches(2.6), Inches(3.5), ACCENT if i == 2 else RULE, 19000)
        text(s, x, Inches(2.85), Inches(3.5), Inches(3.5), body, 17, INK, line=1.25)
    note(s, LM, Inches(6.2), Inches(11), "Full review with sources was presented in Review I.", 13)

    # 6 pain points -> design
    s = blank(prs, R, 6, title="Pain points and how the design answers them", section="Why it is built this way")
    rows = [["Billing data arrives 8 to 24 hours late", "Verify only after 3 billed days; waste findings do not wait for the bill"],
            ["Resources are created without tags", "Creator read from the Activity Log (earliest successful write), tag as fallback"],
            ["A long list gives no order", "One score: money, severity, confidence and age; bands P1 to P4"],
            ["Tools ask for broad write access", "Reader and Cost Management Reader only; the app never writes to the cloud"],
            ["Savings are claimed, not proven", "Before / after bill comparison; reopen and withdraw if the resource returns"],
            ["Failures are silent", "Per-step errors, a check-azure command and a start-up message that says why"]]
    table(s, LM, Inches(2.0), [4.4, 7.2], ["Pain point in practice", "Design answer"], rows, size=14, row_h=0.72, bold_first=True)

    architecture_slide(prs, R, 7)

    # 8 credential flow
    s = blank(prs, R, 8, title="How CloudPulse gets into Azure", section="Azure credential architecture")
    pic(s, A + "fig_credflow.png", Inches(0.85), Inches(1.65), h=Inches(5.2), border=False)
    rich(s, Inches(5.5), Inches(1.95), Inches(6.9), Inches(4.9), [
        ("Identity.", "An app registration with a service principal in Microsoft Entra ID. Three values: tenant ID, client ID, client secret."),
        ("Permission.", "Reader and Cost Management Reader on the subscription only. No role can create, change or delete."),
        ("Token.", "Client-credentials request returns a JWT valid about an hour; sent as a Bearer token on every read."),
        ("Storage.", "A secret entered in the UI is encrypted with Fernet and never returned by the API."),
        ("Alternatives.", "Managed identity on App Service (no secret at all) or an az CLI login on a laptop."),
        ("Diagnosis.", "python -m app.manage check-azure prints Connected or the exact reason it failed.")], 15, 8)

    # 9 modules 1
    s = blank(prs, R, 9, title="Module description (1 of 3): access and data", section="Modules")
    mods = [("A  Azure access", "connections.py, azure_rest.py",
             "Builds a credential from stored or server settings; custom ARM REST client with retry on 429 / 5xx, Retry-After and $skipToken paging."),
            ("B  Data collection", "azure_collector.py, collection_orchestrator.py",
             "Cost Management query, Resource Graph inventory, Monitor metrics (14 days), Activity Log, Advisor, Retail Prices. Optional AWS and FOCUS import."),
            ("C  Ownership resolver", "owners.py",
             "Earliest successful write in the Activity Log; GUID callers shown as service identities; then owner tags; cached 24 h when unknown.")]
    for i, (t, f, d) in enumerate(mods):
        y = Inches(1.95) + Inches(1.6) * i
        text(s, LM, y, Inches(3.3), Inches(0.5), t, 21, INK, True, font=HEAD)
        text(s, LM, y + Inches(0.5), Inches(3.3), Inches(0.4), f, 12, MUTED, italic=True)
        text(s, Inches(4.5), y + Inches(0.02), Inches(7.9), Inches(1.3), d, 17, INK, line=1.2)
        rule(s, LM, y + Inches(1.38), CW)

    # 10 modules 2
    s = blank(prs, R, 10, title="Module description (2 of 3): detection and analytics", section="Modules")
    rows = [["D  Waste detectors", "Unattached disk, static IP on nothing, orphaned snapshot, stopped-but-billed VM", "Rules on inventory; confidence 0.90 to 0.95"],
            ["E  VM verdict", "Idle: p95 CPU < 5 % and network < 5 MB / h. Oversized: p95 CPU < 20 %, peak < 80 %", "14 days of hourly metrics"],
            ["F  Anomaly detector", "z = (x − median) / (1.4826 × MAD) > 3.5 against the same weekday of 4 weeks", "Daily cost series"],
            ["G  Forecast", "Holt smoothing + weekday factors; 80 % range; first day the total passes the budget", "Month-to-date cost"],
            ["H  Priority score", "100 × (0.6 × impact + 0.4 × severity) × confidence × age; P1 ≥ 70, P2 ≥ 45, P3 ≥ 25", "Finding evidence"]]
    table(s, LM, Inches(2.0), [2.3, 6.2, 3.1], ["Module", "What it does", "Input"], rows, size=13, row_h=0.82, bold_first=True)

    # 11 modules 3
    s = blank(prs, R, 11, title="Module description (3 of 3): findings, security and interface", section="Modules")
    pic(s, A + "fig_lifecycle.png", Inches(0.85), Inches(1.95), w=Inches(5.7), border=False)
    caption(s, Inches(0.85), Inches(5.8), Inches(6), "Finding lifecycle: open, assigned, in progress, resolved, verified.")
    rich(s, Inches(7.3), Inches(1.95), Inches(5.2), Inches(4.9), [
        ("I  Findings engine.", "Fingerprint de-duplication, scoring, lifecycle, 7-day verification, savings ledger."),
        ("J  API and security.", "scrypt passwords, server-side sessions, CSRF header, login rate limit, role checks, audit log."),
        ("K  Web interface.", "Ten pages with SVG charts: overview, costs, anomalies, budgets, findings, owners, savings, sources, admin, audit."),
        ("L  Azure scripts.", "Create workload, plant four leaks, read-only identity, deploy, cost spike, clean-up.")], 14, 9)
    rich(s, Inches(0.85), Inches(6.15), Inches(6.2), Inches(0.7), [
        ("Roles:", "admin and owner can do everything; the assignee can start and resolve; viewers read.")], 13, 0)

    # 12 reports
    s = blank(prs, R, 12, title="What the system reports", section="Outputs")
    rows = [["Overview", "Spend, findings to fix, verified savings, loop by stage"],
            ["Cost explorer", "Spend by service or provider with a period comparison"],
            ["Anomalies", "Flagged days, excess, drivers, ongoing or one-off"],
            ["Budgets", "Cumulative spend, forecast with 80 % range, breach date"],
            ["Findings and owners", "Ranked inbox, filters, CSV export; who owns what"],
            ["Savings ledger", "Estimate against realised saving, accuracy, saved to date"],
            ["Audit log", "Every sign-in, change and action, newest first"]]
    table(s, LM, Inches(1.95), [2.8, 8.8], ["Report", "What it shows"], rows, size=14, row_h=0.52, bold_first=True)
    note(s, LM, Inches(6.55), Inches(11), "Next, Review III: results from the running system, testing, conclusion and future enhancements.", 14)
    prs.save(OUT / "PD_Lab_III_Review_2.pptx")


# ====================================================================== REVIEW 3
def review3():
    prs = new_deck(); R = 3
    title_slide(prs, R, "Review III")
    agenda(prs, R, [("Abstract", 3), ("Objectives and outcome", 4), ("Literature review", 5), ("System architecture", 6),
                    ("Module description", 7), ("Results", "8 – 9"), ("Testing and validation", 10), ("Conclusion", 11),
                    ("Future enhancements", 12)])
    abstract_slide(prs, R, 3)

    # 4 objectives with outcome
    s = blank(prs, R, 4, title="Objectives and what we achieved", section="Outcome against plan")
    rows = [["Collect read-only Azure data", "Service principal, CLI or managed identity; Cost, Graph, Monitor, Activity Log, Advisor"],
            ["Detect waste", "Four resource rules plus idle / oversized VM verdict; 12 findings in the demo"],
            ["Explain spend", "24 Sep flagged ($169.26 vs usual $73.55, z = 7.9); budget breach on 28 Sep"],
            ["Assign owners", "Creator from the Activity Log, tags second; 6 of 7 idle resources traced"],
            ["Verify savings", "2 findings verified: $33.44 realised against $35.71 estimated"],
            ["Prove it", "27 automated tests pass; Azure scripts plant four leaks"]]
    table(s, LM, Inches(2.0), [3.3, 8.3], ["Objective", "Outcome"], rows, size=14, row_h=0.64, bold_first=True)

    # 5 literature
    s = blank(prs, R, 5, title="Literature review and the gap", section="Background")
    for i, (h, body) in enumerate([
            ("What exists", "Consoles report, platforms allocate behind closed scoring, scripts give one-off lists [3], [4]."),
            ("What we used", "Robust z-score [5], [6]; Holt smoothing [7], [8]; the FinOps loop [3], [9]; OAuth 2.0 and RBAC [10], [11]."),
            ("What we added", "Owner from the audit log, one explainable priority and savings verified on the bill, read-only.")]):
        x = LM + Inches(3.95) * i
        text(s, x, Inches(2.0), Inches(3.6), Inches(0.5), h, 22, ACCENT if i == 2 else INK, True, font=HEAD)
        rule(s, x, Inches(2.6), Inches(3.5), ACCENT if i == 2 else RULE, 19000)
        text(s, x, Inches(2.85), Inches(3.5), Inches(3.5), body, 17, INK, line=1.25)

    architecture_slide(prs, R, 6)

    # 7 modules table
    s = blank(prs, R, 7, title="Module description", section="Design")
    rows = [["Access", "connections, azure_rest", "Credential, token, retry, paging"],
            ["Collection", "azure_collector, orchestrator", "Bill, inventory, metrics, advice, audit log"],
            ["Ownership", "owners", "Activity Log, tags, cache"],
            ["Detection", "azure_collector, analytics", "Waste rules, VM verdict"],
            ["Analytics", "analytics", "Anomaly, forecast, priority score"],
            ["Findings", "findings", "Fingerprint, lifecycle, verification, ledger"],
            ["API and security", "routers, security, audit", "Sessions, roles, CSRF, audit trail"],
            ["Interface", "static/app.js, charts.js", "Ten pages, SVG charts, light and dark themes"]]
    table(s, LM, Inches(1.95), [2.3, 3.6, 5.7], ["Module", "Files", "Responsibility"], rows, size=14, row_h=0.56, bold_first=True)

    # 8 results screenshots
    s = blank(prs, R, 8, title="Results: the running system", section="Results 1 / 2")
    shots = [("shot_overview.png", "Overview: $2,829.99 spend; $262.94 to fix; $33.44 verified"),
             ("shot_costs.png", "Cost explorer: 80 services, daily average $94.33"),
             ("shot_anomalies.png", "Anomalies: 24 Sep, excess $95.71, 99 % from EC2"),
             ("shot_budgets.png", "Budgets: $2,600 passed on 28 Sep, $229.99 over")]
    for i, (f, c) in enumerate(shots):
        x = LM + Inches(4.55) * (i % 2); y = Inches(1.85) + Inches(2.6) * (i // 2)
        pic(s, A + f, x, y, w=Inches(4.3))
        caption(s, x, y + Inches(2.25), Inches(4.3), c)
    text(s, Inches(10.0), Inches(1.9), Inches(2.5), Inches(0.4), "What to notice", 17, INK, True, font=HEAD)
    dash_list(s, Inches(10.0), Inches(2.4), Inches(2.45), Inches(4.2), [
        "The same totals appear on every page",
        "Tooltips go down to one day and one service",
        "The budget breach is marked on the chart, not only in a number"], 14, 10)

    # 9 results numbers
    s = blank(prs, R, 9, title="Results: findings and verified savings", section="Results 2 / 2")
    rows = [["1", "September 2024 ended 9 % over budget", "230.00", "75  P1", "Open"],
            ["2", "VM gpu-trial-01 stopped but still billed", "70.08", "69  P2", "In progress"],
            ["3", "VM reporting-vm runs but does almost nothing", "70.08", "65  P2", "Open"],
            ["4", "VM api-prod-01 larger than it needs to be", "70.08", "54  P2", "Open"],
            ["5", "Disk web-01_OsDisk_old attached to nothing", "19.71", "43  P3", "Verified"],
            ["6", "EBS volume db-backup-old attached to nothing", "16.00", "39  P3", "Verified"]]
    table(s, LM, Inches(1.95), [0.5, 4.3, 1.1, 1.1, 1.2], ["#", "Finding (top of 12)", "USD / mo", "Score", "Status"], rows, size=12, row_h=0.5)
    panel(s, Inches(9.35), Inches(1.95), Inches(3.1), Inches(3.5))
    text(s, Inches(9.6), Inches(2.15), Inches(2.7), Inches(0.3), "VERIFIED SAVINGS", 11, ACCENT, True, spacing=2)
    text(s, Inches(9.6), Inches(2.55), Inches(2.7), Inches(0.7), "$33.44", 36, INK, True, font=HEAD)
    text(s, Inches(9.6), Inches(3.3), Inches(2.7), Inches(1.0), "a month, from an estimate of $35.71", 13, MUTED, line=1.2)
    text(s, Inches(9.6), Inches(4.1), Inches(2.7), Inches(0.7), "93.6 %", 28, INK, True, font=HEAD)
    text(s, Inches(9.6), Inches(4.75), Inches(2.7), Inches(0.5), "estimate accuracy", 13, MUTED)
    note(s, LM, Inches(5.75), Inches(8.2), "All figures are from the demo workspace on the FOCUS sample bill (21 Mar to 30 Sep 2024), not a customer account.", 13)
    note(s, LM, Inches(6.15), Inches(8.2), "Verification showed estimates are slightly optimistic, which a plain report would never reveal.", 13)

    # 10 testing
    s = blank(prs, R, 10, title="Testing and validation", section="Evidence")
    rows = [["Algorithms", "23", "Median / MAD, anomaly, forecast, priority, VM verdict, before / after"],
            ["Full loop", "1", "Fake Azure: sync, owners, workflow, verify, reappearing resource, permissions, audit"],
            ["Front end and diagnostics", "3", "CSS colour tokens, page renderers, Azure setup messages"]]
    table(s, LM, Inches(2.0), [3.0, 0.9, 4.2], ["Suite", "Tests", "What is checked"], rows, size=13, row_h=0.72, bold_first=True)
    text(s, Inches(9.4), Inches(2.0), Inches(3.1), Inches(0.4), "Planted leaks (live check)", 16, INK, True, font=HEAD)
    dash_list(s, Inches(9.4), Inches(2.6), Inches(3.1), Inches(3), [
        "leftover-disk, 4 GB", "unused-ip, static", "old-backup-snap, disk deleted", "forgotten-vm, stopped not deallocated"], 14, 7)
    rich(s, LM, Inches(5.3), Inches(8.1), Inches(1.6), [
        ("Found by a browser review:", "a mistyped colour value made the main text nearly invisible. Fixed, and a regression test was added."),
        ("Not yet measured:", "results on a customer subscription; Azure bills arrive 8 to 24 hours late.")], 14, 8)

    # 11 conclusion
    s = blank(prs, R, 11, title="Conclusion", section="Wrap-up")
    text(s, LM, Inches(2.0), Inches(7.2), Inches(2.6),
         "CloudPulse carries a cost problem from the first signal to a saving proven on the bill. It reads Azure with a read-only identity, finds waste and "
         "anomalies, forecasts budget breaches, names the owner, tracks the fix and withdraws the saving if the problem returns.", 19, INK, line=1.25)
    dash_list(s, LM, Inches(4.35), Inches(7.2), Inches(2.4), [
        "Every figure is explainable and recomputable from stored evidence.",
        "27 automated tests and a scripted Azure environment back the results.",
        "Known limits: billing delay, no VM memory metrics, creators older than 89 days only by tag."], 16, 9)
    panel(s, Inches(8.5), Inches(2.0), Inches(3.95), Inches(4.4))
    text(s, Inches(8.8), Inches(2.25), Inches(3.4), Inches(0.3), "WHAT WE TAKE AWAY", 11, ACCENT, True, spacing=2)
    dash_list(s, Inches(8.8), Inches(2.75), Inches(3.4), Inches(3.5), [
        "Cloud billing, IAM and monitoring applied in one product", "Explainable algorithms are easier to defend", "Proof beats estimates"], 15, 12)

    # 12 future
    s = blank(prs, R, 12, title="Future enhancements", section="What comes next")
    rows = [["Sign in with Microsoft Entra ID", "Use existing organisation accounts"],
            ["E-mail, Teams and Slack alerts", "Reach owners outside the app"],
            ["Google Cloud and deeper AWS", "Idle load balancers, NAT gateways"],
            ["Reservation and savings-plan analysis", "Purchase recommendations"],
            ["Tag-compliance policies, showback", "Cost by team"],
            ["Approval-gated remediation", "Deallocate a stopped VM after the owner confirms"],
            ["PostgreSQL default, shared job queue", "Multi-instance deployment"],
            ["VM memory and disk metrics", "More accurate right-sizing"]]
    table(s, LM, Inches(1.95), [5.0, 6.6], ["Enhancement", "Why"], rows, size=14, row_h=0.5, bold_first=True)
    text(s, LM, Inches(6.5), Inches(11), Inches(0.4), "Thank you. Questions are welcome.", 16, INK, italic=True, font=HEAD)
    prs.save(OUT / "PD_Lab_III_Review_3.pptx")


if __name__ == "__main__":
    review1(); review2(); review3()
    print("ok")
