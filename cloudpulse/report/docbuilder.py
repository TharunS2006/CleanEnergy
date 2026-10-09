"""Small python-docx layer that reproduces the RMKCET mini-project template:
A4, Times New Roman 12, boxed page border, front matter without page numbers,
body numbering restarting at 1."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def _set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _page_border(section):
    sect = section._sectPr
    for old in sect.findall(qn("w:pgBorders")):
        sect.remove(old)
    b = OxmlElement("w:pgBorders")
    b.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "24")
        e.set(qn("w:color"), "000000")
        b.append(e)
    # pgBorders must come after pgMar in sectPr
    pgmar = sect.find(qn("w:pgMar"))
    pgmar.addnext(b)


def _field(paragraph, instr, cached="1", size=11):
    def mk(t, **attrs):
        e = OxmlElement(t)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r1 = paragraph.add_run()
    r1._element.append(mk("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = paragraph.add_run()
    it = mk("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._element.append(it)
    r3 = paragraph.add_run()
    r3._element.append(mk("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = paragraph.add_run(cached)
    r5 = paragraph.add_run()
    r5._element.append(mk("w:fldChar", **{"w:fldCharType": "end"}))
    for r in (r1, r2, r3, r4, r5):
        _set_font(r, size)


class Report:
    cm = staticmethod(Cm)

    def __init__(self):
        self.captions = []
        self.doc = Document()
        self.fig_no = {}
        d = self.doc
        st = d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        for name, size, before, after in (("Heading 1", 14, 6, 12), ("Heading 2", 12.5, 12, 4), ("Heading 3", 12, 8, 2)):
            h = d.styles[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = name == "Heading 3"
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
            h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15
        for name in ("List Bullet",):
            b = d.styles[name]
            b.font.name = FONT
            b.font.size = Pt(12)
            b.paragraph_format.line_spacing = 1.5
            b.paragraph_format.space_after = Pt(3)
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin = Cm(3.0), Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.3)
        s.footer_distance = Cm(1.2)
        _page_border(s)
        self.front = s
        self.body = None

    # ------------------------------------------------------------------ text
    def _runs(self, p, text, size=12, base_bold=False, base_italic=False):
        for tok in re.split(r"(\*\*.+?\*\*|`.+?`|__.+?__|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))", text):
            if not tok:
                continue
            if tok.startswith("**"):
                r = p.add_run(tok[2:-2]); _set_font(r, size, True, base_italic)
            elif tok.startswith("`"):
                r = p.add_run(tok[1:-1]); _set_font(r, size - 1, base_bold, base_italic, name="Courier New")
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = p.add_run(tok[1:-1]); _set_font(r, size, base_bold, True)
            elif tok.startswith("__"):
                r = p.add_run(tok[2:-2]); _set_font(r, size, base_bold, True)
            else:
                r = p.add_run(tok); _set_font(r, size, base_bold, base_italic)

    def para(self, text, align="j", size=12, bold=False, italic=False, before=0, after=6, spacing=1.5,
             indent=None, keep=False):
        p = self.doc.add_paragraph()
        p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                       "l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = spacing
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
        if keep:
            p.paragraph_format.keep_with_next = True
        self._runs(p, text, size, bold, italic)
        return p

    def h(self, level, text):
        p = self.doc.add_heading(level=level)
        r = p.add_run(text)
        _set_font(r, {1: 14, 2: 12.5, 3: 12}[level], True, level == 3)
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def bullets(self, items):
        for t in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs(p, t)

    def numbered(self, items, prefix=""):
        for i, t in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(3)
            self._runs(p, f"{prefix}{i}.  {t}")

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Cm(0.6)
            pPr = p._element.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
            pPr.append(shd)
            r = p.add_run(line if line else " ")
            _set_font(r, 9, name="Courier New")
        self.para("", after=2, spacing=1.0)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ----------------------------------------------------------------- table
    def table(self, caption, header, rows, widths=None, size=10.5, align_first="l"):
        self.captions.append(caption)
        cap = self.para(caption, align="c", size=11, bold=True, before=6, after=4, spacing=1.0, keep=True)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            self._runs(p, htxt, size, True)
            _shade(c, "D9D9D9")
        # repeat header row
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._runs(p, str(txt), size)
        for r in t.rows:
            cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true")
            r._tr.get_or_add_trPr().append(cant)
        if widths:
            self._fix_widths(t, widths)
        if len(rows) <= 14:
            for row in t.rows[:-1]:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.paragraph_format.keep_with_next = True
        self.para("", after=4, spacing=1.0)
        return t

    @staticmethod
    def _fix_widths(t, widths):
        t.autofit = False
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblLayout")):
            tblPr.remove(old)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 567))); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    # ---------------------------------------------------------------- figure
    def figure(self, path, caption, width_cm=14.0, keep_caption=True):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = keep_caption
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.captions.append(caption)
        self.para(caption, align="c", size=11, bold=True, after=10, spacing=1.0)

    # -------------------------------------------------------------- sections
    def start_body(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.footer.is_linked_to_previous = False
        s.header.is_linked_to_previous = False
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
        s._sectPr.append(pg)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _field(fp, "PAGE", "1")
        self.body = s
        return s

    def landscape_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.8)
        for e in s._sectPr.findall(qn("w:pgNumType")):
            s._sectPr.remove(e)
        _page_border(s)
        return s

    def portrait_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin = Cm(3.0), Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.3)
        for e in s._sectPr.findall(qn("w:pgNumType")):
            s._sectPr.remove(e)
        _page_border(s)
        return s

    def save(self, path):
        for s in self.doc.sections:
            _page_border(s)
        self.doc.save(path)
