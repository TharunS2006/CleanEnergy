"""Builds the internship report (DOCX + PDF) in the layout of the RMKCET internship-report reference."""
import re, subprocess, sys, pathlib
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from ibuilder import IReport, _font, fix_widths, NAVY

HERE = pathlib.Path(__file__).parent
A = str(HERE / "assets") + "/"

ST = dict(name="TAMIL MANI B", reg="111624102217", batch="2024 to 2028",
          org="Trios Technologies Private Limited", dur="08 June 2026 to 22 June 2026",
          guide="Ms. Gokila E & Dr. Uma Mageshwari G", mentor="Ms. Kalaiyarasi S", year="2026-2027",
          addr="Rani Arcade, #24/5, 1st Avenue, Ashok Nagar, Chennai - 83. (Near Ashok Pillar)")

ABSTRACT = (
    "This report describes the practical knowledge and experience I gained during a 15-day internship at Trios Technologies "
    "Private Limited, Ashok Nagar, Chennai, from 08 June to 22 June 2026. The internship was in the domain of cloud computing and "
    "gave me exposure to cloud service models, virtualisation and compute, storage, networking, identity and access management, "
    "resource governance, monitoring, cloud pricing and FinOps. I applied these concepts to a mini project, CloudPulse, whose "
    "development was started during the internship and extended afterwards. CloudPulse is a cloud cost monitoring and optimisation "
    "platform for Microsoft Azure. It connects with a read-only service principal, reads the bill, resource inventory, virtual "
    "machine metrics and the Activity Log, detects idle disks, unused public IP addresses, orphaned snapshots and stopped-but-billed "
    "virtual machines, finds spend anomalies and forecasts budget breaches, assigns every finding to the person who created the "
    "resource, and counts a saving only after seven days of billed cost confirm it. The tools used included Python, FastAPI, "
    "SQLAlchemy, JavaScript, the Azure CLI and Azure Resource Manager APIs, Docker, Git and GitHub. The project is validated by 27 "
    "automated tests and, in its demonstration workspace, reported $262.94 a month of open waste and $33.44 a month of verified "
    "savings. Through these activities I improved my cloud computing knowledge, problem-solving, security awareness, testing and "
    "documentation skills, and gained a clearer view of how cloud applications are designed, secured and operated.")

ABBREV = [("API", "Application Programming Interface"), ("ARM", "Azure Resource Manager"), ("AWS", "Amazon Web Services"),
          ("CLI", "Command-Line Interface"), ("CSS", "Cascading Style Sheets"), ("CSV", "Comma-Separated Values"),
          ("FinOps", "Cloud Financial Operations"), ("FOCUS", "FinOps Open Cost and Usage Specification"),
          ("HTML", "HyperText Markup Language"), ("HTTP", "HyperText Transfer Protocol"), ("IaaS", "Infrastructure as a Service"),
          ("IAM", "Identity and Access Management"), ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"),
          ("OBE", "Outcome Based Education"), ("ORM", "Object-Relational Mapper"), ("PaaS", "Platform as a Service"),
          ("PO", "Program Outcome"), ("PSO", "Program Specific Outcome"), ("RBAC", "Role-Based Access Control"),
          ("REST", "Representational State Transfer"), ("SaaS", "Software as a Service"), ("SDG", "Sustainable Development Goal"),
          ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("SVG", "Scalable Vector Graphics"),
          ("UI", "User Interface"), ("VM", "Virtual Machine")]


class R(IReport):
    def __init__(self):
        super().__init__()
        self.heads = []

    def h(self, level, text):
        self.heads.append((level, text))
        return super().h(level, text)


def cover(r):
    t = r.doc.add_table(rows=1, cols=2)
    fix_widths(t, [3.0, 12.9])
    t.rows[0].cells[0].paragraphs[0].add_run().add_picture(A + "crest_slide.png", width=Cm(2.3))
    c = t.rows[0].cells[1]; c.text = ""
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10)
    _font(p.add_run("R.M.K. COLLEGE OF ENGINEERING AND TECHNOLOGY"), 13, True, color=NAVY)
    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("Department of Computer Science and Engineering"), 11, True)
    r.para("", after=24)
    r.para("INTERNSHIP REPORT", "c", 20, True, color=NAVY, after=10)
    r.para("Submitted in partial fulfilment of the requirements for the award of the degree of", "c", 11, after=6, spacing=1.15)
    r.para("Bachelor of Engineering", "c", 13, True, after=2)
    r.para("in", "c", 11, after=2)
    r.para("Computer Science and Engineering", "c", 12, after=2)
    r.para("24CS312 - Internship & Seminar", "c", 11.5, True, after=14)
    r.para("By", "c", 12, after=14)
    for k, v in (("Name of the Student", ST["name"]), ("Register Number", ST["reg"]), ("Batch", ST["batch"]),
                 ("Internship Organization", ST["org"]), ("Internship Duration", ST["dur"]),
                 ("Faculty Guide", ST["guide"]), ("Industry Mentor", ST["mentor"]), ("Academic Year", ST["year"])):
        r.para(f"{k} :  {v}", "l", 11.5, after=4, spacing=1.15, indent=1.0)
    r.page_break()


def bonafide(r):
    r.para("BONAFIDE CERTIFICATE", "c", 14, True, color=NAVY, after=18)
    r.para("This is to certify that the ‘Internship’ report submitted by", "c", after=14)
    r.para(f"Name: **{ST['name']}**", "l", indent=2.0, after=4)
    r.para(f"Reg. No: **{ST['reg']}**", "l", indent=2.0, after=14)
    r.para(f"is work completed by {ST['name'].title().replace('B', 'B')} and submitted during the academic year {ST['year']}, in partial "
           "fulfilment of the requirements for the award of the degree of Bachelor of Engineering in the Department of Computer Science "
           "and Engineering at R.M.K. College of Engineering and Technology. The internship work was carried out at "
           f"{ST['org']}, Ashok Nagar, Chennai.", spacing=1.5)
    r.para("", after=40)
    t = r.doc.add_table(rows=1, cols=2)
    fix_widths(t, [7.9, 8.0])
    for cell, lines in zip(t.rows[0].cells, (("Faculty In-charge", "Signature: ____________________", "Department of CSE"),
                                              ("Head of the Department", "Dr. R. Raja Guru", "Department of CSE"))):
        cell.text = ""
        for i, ln in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
            _font(p.add_run(ln), 12, i == 0)
    r.para("", after=18)
    r.para("Course Coordinator (24CS312)", "l", 12, True, after=2)
    r.para("Mr. K. Suresh Babu", "l", 12, after=2, spacing=1.0)
    r.para("Department of CSE", "l", 12, after=22, spacing=1.0)
    r.para("Place: ____________________", "l", after=2)
    r.para("Date: ____________________", "l")
    r.page_break()


def certificate(r):
    r.para("INTERNSHIP CERTIFICATE", "c", 14, True, color=NAVY, after=8)
    r.figure(A + "certificate.png", f"Internship completion letter issued by {ST['org']}", width_cm=13.6)
    r.page_break()


def declaration(r):
    r.para("STUDENT DECLARATION", "c", 14, True, color=NAVY, after=18)
    r.para("I hereby declare that this internship report is my original work carried out during the internship period under the "
           "guidance of the industry mentor and faculty guide. The work presented in this report has not been submitted elsewhere "
           "for any other academic award.", spacing=1.5)
    r.para("", after=40)
    r.para(ST["name"].title().replace("B", "B"), "r", 12, True, after=4)
    r.para("Date: ____________________", "r")
    r.page_break()


def acknowledgement(r):
    r.para("ACKNOWLEDGEMENT", "c", 14, True, color=NAVY, after=14)
    for t in (
        f"I would like to thank Ms. Kalaiyarasi S, my industry mentor at {ST['org']}, for providing me the opportunity to undertake "
        "internship training in the domain of cloud computing and for her guidance during the internship. The exposure to cloud "
        "computing concepts helped me connect classroom learning with the way cloud services are designed, secured and paid for.",
        f"I am grateful to the people at {ST['org']} for their support and for the learning environment they provided.",
        "I also thank the management and Principal Dr. N. Suresh Kumar for providing the facilities needed to complete this "
        "internship.",
        "I extend my thanks to Dr. R. Raja Guru, Professor and Head of the Department of Computer Science and Engineering, for "
        "permitting me to undertake the internship, to my faculty guides Ms. Gokila E and Dr. Uma Mageshwari G, to Mr. K. Suresh "
        "Babu, the department internship coordinator, and to the faculty members and friends for their support during its completion."):
        r.para(t, spacing=1.5)
    r.para("", after=10)
    r.para(ST["name"].title().replace("B", "B"), "r", 12, True)
    r.page_break()


def abstract(r):
    r.para("ABSTRACT", "c", 14, True, color=NAVY, after=14)
    r.para(ABSTRACT, spacing=1.5)
    r.page_break()


def toc(r, pages):
    r.para("TABLE OF CONTENTS", "c", 14, True, color=NAVY, after=10)
    t = r.doc.add_table(rows=1, cols=2); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Section / Content", "Page")):
        c.text = ""; q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3); q.paragraph_format.line_spacing = 1.0
        _font(q.add_run(h), 11, True)
    for level, text in pages["entries"]:
        cells = t.add_row().cells
        for c, txt in zip(cells, (text, str(pages["of"].get(text, "")))):
            c.text = ""; q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(1.5); q.paragraph_format.space_before = Pt(1.5); q.paragraph_format.line_spacing = 1.0
            if level == 1:
                _font(q.add_run(txt), 10.5, True)
            else:
                q.paragraph_format.left_indent = Cm(0.0)
                _font(q.add_run(("    " if level == 3 else "") + txt), 10.5)
    fix_widths(t, [12.0, 3.9])
    r.page_break()


def lists(r, pages):
    for title, items, key in (("LIST OF FIGURES", pages["figs"], "Figure"), ("LIST OF TABLES", pages["tabs"], "Table")):
        r.para(title, "c", 14, True, color=NAVY, after=10)
        t = r.doc.add_table(rows=1, cols=3); t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, (f"{key} No.", f"{key} Title", "Page No.")):
            c.text = ""; q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3); q.paragraph_format.line_spacing = 1.0
            _font(q.add_run(h), 11, True)
        for cap in items:
            m = re.match(rf"{key} (\d+) - (.*)", cap)
            cells = t.add_row().cells
            for i, (c, txt) in enumerate(zip(cells, (m.group(1), m.group(2), str(pages["of"].get(cap, ""))))):
                c.text = ""; q = c.paragraphs[0]
                q.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.space_after = Pt(2.5); q.paragraph_format.space_before = Pt(2.5); q.paragraph_format.line_spacing = 1.0
                _font(q.add_run(txt), 10.5)
        fix_widths(t, [2.2, 11.5, 2.2])
        r.page_break()
    r.para("LIST OF ABBREVIATIONS", "c", 14, True, color=NAVY, after=10)
    t = r.doc.add_table(rows=1, cols=2); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("Abbreviation", "Full Form")):
        c.text = ""; q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3); q.paragraph_format.line_spacing = 1.0
        _font(q.add_run(h), 11, True)
    for a, f in ABBREV:
        cells = t.add_row().cells
        for i, (c, txt) in enumerate(zip(cells, (a, f))):
            c.text = ""; q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            q.paragraph_format.space_after = Pt(2); q.paragraph_format.space_before = Pt(2); q.paragraph_format.line_spacing = 1.0
            _font(q.add_run(txt), 10.5)
    fix_widths(t, [3.5, 12.4])
    r.page_break()


from internship_content import chapters  # noqa: E402


def build(out, pages):
    r = R()
    cover(r); bonafide(r); certificate(r); declaration(r); acknowledgement(r); abstract(r)
    toc(r, pages); lists(r, pages)
    chapters(r)
    r.save(out)
    return r


def pdf(docx, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)


def texts(pdfp):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdfp)], capture_output=True, text=True).stdout).group(1))
    return [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdfp), "-"], capture_output=True, text=True).stdout
            for i in range(1, n + 1)]


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def locate(pdfp, entries, figs, tabs):
    tx = texts(pdfp)
    lines = [[norm(l) for l in t.splitlines()] for t in tx]
    abbr = next(i for i, ls in enumerate(lines) if "LIST OF ABBREVIATIONS" in ls)
    start = next(i for i in range(abbr + 1, len(lines)) if "CHAPTER 1 - INTRODUCTION" in lines[i])
    of, cur = {}, start
    for level, text in entries:
        for i in range(cur, len(lines)):
            if norm(text) in lines[i]:
                of[text] = i + 1; cur = i; break
    for cap in list(figs) + list(tabs):
        key = cap.split(" - ")[0] + " - "
        for i in range(len(lines)):
            if any(l.startswith(key) for l in lines[i]):
                of[cap] = i + 1; break
    return {"of": of, "total": len(tx)}


if __name__ == "__main__":
    out = HERE / "output"; out.mkdir(exist_ok=True)
    docx = out / "Internship_Report_Tamil_Mani_B.docx"
    empty = {"entries": [], "of": {}, "figs": [], "tabs": []}
    r = build(docx, empty)                                    # pass 0: collect headings and captions
    entries = [(l, t) for l, t in r.heads]
    figs, tabs = list(r.figures), list(r.tables)
    pages = {"entries": entries, "of": {}, "figs": figs, "tabs": tabs}
    for _ in range(3):
        build(docx, pages); pdf(docx, out)
        loc = locate(out / docx.with_suffix(".pdf").name, entries, figs, tabs)
        if loc["of"] == pages["of"]:
            break
        pages["of"] = loc["of"]
    print("pages:", loc["total"], "| located", len(loc["of"]), "of", len(entries) + len(figs) + len(tabs))
    missing = [t for _, t in entries if t not in loc["of"]] + [c for c in figs + tabs if c not in loc["of"]]
    print("missing:", missing)
