"""Builds the internship review presentation (.pptx) in the layout of the RMKCET reference deck:
college header, yellow title bar with red title, bullet body, 4:3 slides."""
import pathlib
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).parent
A = str(HERE / "assets") + "/"
FONT = "Times New Roman"
RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLACK = RGBColor(0, 0, 0)
YELLOW = RGBColor(0xFF, 0xFF, 0x00)
PH = RGBColor(0xC0, 0x00, 0xC0)       # placeholder colour (magenta) so it cannot be missed

W, H = Inches(10.667), Inches(8.0)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def tb(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    return s, tf


def run(p, text, size=20, bold=False, italic=False, color=BLACK):
    r = p.add_run(); r.text = text
    f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return r


def header(slide, title):
    slide.shapes.add_picture(A + "crest_slide.png", Inches(0.25), Inches(0.12), height=Inches(0.95))
    slide.shapes.add_picture(A + "logo31.png", Inches(9.55), Inches(0.12), height=Inches(0.9))
    _, tf = tb(slide, Inches(1.35), Inches(0.2), Inches(8.1), Inches(0.9))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "R.M.K. COLLEGE OF ENGINEERING AND TECHNOLOGY", 17, True)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    run(p, "(An Autonomous Institution)", 9, True)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    run(p, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", 11, True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.25), W, Inches(0.62))
    bar.fill.solid(); bar.fill.fore_color.rgb = YELLOW; bar.line.fill.background()
    tf = bar.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, title, 20, True, color=RED)


def bullets(slide, items, x=Inches(0.55), y=Inches(2.15), w=Inches(9.6), h=Inches(5.5), size=21, gap=9):
    _, tf = tb(slide, x, y, w, h)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(gap)
        lead, text, ph = (it + (False,))[:3] if isinstance(it, tuple) else (None, it, False)
        run(p, "•  ", size, True)
        if lead:
            run(p, lead + " ", size, True, color=RED)
        if ph:
            run(p, text, size, False, True, PH)
        else:
            run(p, text, size)
    return tf


def slide(title):
    s = prs.slides.add_slide(BLANK)
    header(s, title)
    return s


# ------------------------------------------------------------- 1. details
s = slide("REVIEW OF INTERNSHIP - DATE : __ / 10 / 2026")
rows = [("Name of the Student", "Tamil Mani B"), ("Register Number", "111624102217"), ("Year / Semester", "3/5"),
        ("Academic Year", "2026-2027"), ("Period of Internship", "08/06/2026 to 22/06/2026"),
        ("Name of the Company", "Trios Technologies Private Limited"),
        ("Address of the company", "Rani Arcade, #24/5, 1st Avenue, Ashok Nagar, Chennai - 83. (Near Ashok Pillar)")]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.5), Inches(2.3), Inches(7.3), Inches(4.6)).table
tbl.columns[0].width = Inches(2.6); tbl.columns[1].width = Inches(4.7)
for i, (k, v) in enumerate(rows):
    for j, txt in enumerate((k, v)):
        c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xE7, 0xEE, 0xF6) if j == 0 else RGBColor(0xF5, 0xF8, 0xFB)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        run(tf.paragraphs[0], txt, 15, j == 0, color=RED if j == 0 else BLACK)
ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.2), Inches(2.4), Inches(2.0), Inches(2.4))
ph.fill.background(); ph.line.color.rgb = PH; ph.line.dash_style = 4
tf = ph.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "[Paste your photo here]", 13, False, True, PH)

# ------------------------------------------------------------- 2. contents
s = slide("TABLE OF CONTENTS")
_, tf = tb(s, Inches(1.0), Inches(2.3), Inches(8.8), Inches(5.3))
items = [("Abstract", None), ("About the Company – ", "Trios Technologies Private Limited"), ("Plan of Internship", None),
         ("Training Task and Outcomes", None), ("Project Architecture", None), ("Project Results", None), ("Conclusion", None),
         ("Internship Photographs and Certificate", None)]
for i, (a, b) in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(12)
    run(p, "•  ", 24, True); run(p, a, 24, True, color=NAVY)
    if b: run(p, b, 24, True, color=RED)

# ------------------------------------------------------------- 3. abstract
s = slide("Abstract")
bullets(s, [
    "Completed a 15-day Cloud Computing internship at Trios Technologies Private Limited, Chennai, from 08 June to 22 June 2026.",
    "Learned cloud service models, virtualisation, storage, networking, identity and access management, monitoring and cloud pricing.",
    "Started the mini project CloudPulse: a cloud cost monitoring and optimisation platform for Microsoft Azure.",
    "CloudPulse finds wasted cloud spend, assigns each problem to its creator and proves every saving on the next bill.",
    "Built with Python, FastAPI, SQLAlchemy, JavaScript, Azure CLI, Docker and Git; validated by 27 automated tests.",
    "Improved cloud knowledge, secure design, testing and documentation skills."], size=20)

# ------------------------------------------------------------- 4. company
s = slide("About the Company – Trios Technologies Private Limited")
bullets(s, [
    "Trios Technologies Private Limited (Trios Technologies Pvt. Ltd.) is located at Rani Arcade, Ashok Nagar, Chennai - 83.",
    "Tagline: “We There Wherever The Technology Is”.",
    "Website: www.triostechnology.com  |  E-mail: support@triostechnology.com  |  Phone: 044 4854 0077.",
    "Internship domain: Cloud Computing.",
    "Industry mentor: Ms. Kalaiyarasi S.",
    ("", "[Add 1-2 lines on the company's business and services from its website]", True)], size=21)

# ------------------------------------------------------------- 5. plan
s = slide("Plan of Internship")
bullets(s, [
    ("08 Jun – 14 Jun:", "Orientation; cloud service and deployment models, virtualisation and compute, storage, networking, identity and access management; defining the problem of wasted cloud spend."),
    ("15 Jun – 22 Jun:", "Governance, monitoring, cloud pricing and FinOps; architecture planning and the first version of CloudPulse with read-only Azure access and waste detection."),
    ("Final review:", "Reviewed the work, organised notes and project files, and received the internship completion letter.")], size=22, gap=16)

# ------------------------------------------------------------- 6. tasks
s = slide("Training Task and Outcomes")
bullets(s, [
    ("Project:", "CloudPulse – Cloud Cost Monitoring and Optimisation Platform for Microsoft Azure."),
    ("Features:", "idle disk, IP, snapshot and VM detection; spend anomalies; budget forecast; owner from Activity Log; P1–P4 priority; savings verified on the bill."),
    ("Cloud concepts applied:", "VMs and power states, storage, networking, IAM and RBAC, monitoring, pricing and FinOps, security, automation."),
    ("Outcome:", "working application with 27 automated tests; demo reported $262.94/month open waste and $33.44/month verified savings."),
    ("Skills:", "Python, FastAPI, Azure CLI, REST APIs, Docker, Git, testing.")], size=19, gap=10)

# ------------------------------------------------------------- 7. architecture
s = slide("Project Architecture – CloudPulse")
pic = s.shapes.add_picture(A + "fig_architecture.png", Inches(1.03), Inches(1.95), width=Inches(8.6))
_, tf = tb(s, Inches(0.5), Inches(7.62), Inches(9.7), Inches(0.35))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "Read-only Azure access: service principal (Reader + Cost Management Reader) → Entra ID token → Azure Resource Manager", 12, False, True, NAVY)

# ------------------------------------------------------------- 8. results
s = slide("Project Results")
cells = [("shot_overview.png", "Overview: $2,829.99 spend; $262.94/month to fix; $33.44/month verified"),
         ("shot_costs.png", "Cost explorer: 80 services; daily average $94.33"),
         ("shot_anomalies.png", "Anomalies: 24 Sep flagged ($169.26 vs usual $73.55)"),
         ("shot_budgets.png", "Budgets: $2,600 budget passed on 28 Sep")]
for i, (img, cap) in enumerate(cells):
    x = Inches(0.4 + (i % 2) * 5.15); y = Inches(2.1 + (i // 2) * 2.95)
    s.shapes.add_picture(A + img, x, y, width=Inches(4.95))
    _, tf = tb(s, x, y + Inches(2.52), Inches(4.95), Inches(0.4))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, cap, 11, True, color=NAVY)

# ------------------------------------------------------------- 9. conclusion
s = slide("Conclusion")
bullets(s, [
    "Gained practical exposure to cloud computing concepts during the internship.",
    "Started and developed CloudPulse, a read-only Azure cost monitoring and optimisation platform.",
    "Learned how cloud resources are billed, secured and monitored, and how to prove a saving on the bill.",
    "Improved Python, REST API, testing, documentation and problem-solving skills.",
    "Next: Entra ID sign-in, e-mail and Teams alerts, Google Cloud support and approval-gated remediation.",
    "Thank you to Ms. Kalaiyarasi S, my faculty guides Ms. Gokila E and Dr. Uma Mageshwari G, and Trios Technologies Private Limited for their guidance and support."],
    size=20, gap=10)

# ------------------------------------------------------------- 10. photographs
s = slide("Internship Photographs")
ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.3), Inches(4.6), Inches(3.9))
ph.fill.background(); ph.line.color.rgb = PH; ph.line.dash_style = 4
tf = ph.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "[Paste your workspace / internship photograph here]", 14, False, True, PH)
s.shapes.add_picture(A + "shot_overview.png", Inches(5.4), Inches(2.9), width=Inches(4.8))
for x, cap in ((0.5, "My workspace"), (5.4, "CloudPulse – Overview page")):
    _, tf = tb(s, Inches(x), Inches(6.45), Inches(4.7), Inches(0.5))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, cap, 16, True)

# ------------------------------------------------------------- 11. certificate
s = slide("Internship Certificate")
s.shapes.add_picture(A + "certificate.png", Inches(3.35), Inches(1.95), height=Inches(5.85))

# ------------------------------------------------------------- 12. thank you
s = slide("")
_, tf = tb(s, Inches(1.0), Inches(3.0), Inches(8.6), Inches(2.5), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run(p, "Thank you!", 72, True, True, NAVY)

out = HERE / "output"; out.mkdir(exist_ok=True)
prs.save(out / "Internship_PPT_Tamil_Mani_B.pptx")
print("slides:", len(prs.slides))
