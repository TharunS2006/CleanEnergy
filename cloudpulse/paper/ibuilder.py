"""python-docx helpers that reproduce the look of the RMKCET internship-report reference:
Times New Roman 12, navy headings, italic-blue captions, black-bordered centred tables, centred page numbers."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
NAVY = "1F3864"
CAPBLUE = "2E74B5"


def _font(run, size=None, bold=None, italic=None, color=None, name=FONT, highlight=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if highlight: run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _field(p, instr, cached="1", size=10):
    def mk(t, **kw):
        e = OxmlElement(t)
        for k, v in kw.items(): e.set(qn(k), v)
        return e
    runs = []
    r = p.add_run(); r._element.append(mk("w:fldChar", **{"w:fldCharType": "begin"})); runs.append(r)
    r = p.add_run(); it = mk("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._element.append(it); runs.append(r)
    r = p.add_run(); r._element.append(mk("w:fldChar", **{"w:fldCharType": "separate"})); runs.append(r)
    runs.append(p.add_run(cached))
    r = p.add_run(); r._element.append(mk("w:fldChar", **{"w:fldCharType": "end"})); runs.append(r)
    for r in runs: _font(r, size)


def fix_widths(t, widths):
    t.autofit = False
    pr = t._tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        for old in pr.findall(qn(tag)): pr.remove(old)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 567))); tw.set(qn("w:type"), "dxa"); pr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


class IReport:
    cm = staticmethod(Cm)

    def __init__(self):
        self.doc = Document()
        self.tables, self.figures = [], []
        d = self.doc
        st = d.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 14), ("Heading 2", 12.5), ("Heading 3", 12)):
            h = d.styles[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.italic = False
            h.font.color.rgb = RGBColor.from_string(NAVY)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                h.element.rPr.rFonts.set(qn(a), FONT)
            h.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 0)
            h.paragraph_format.space_after = Pt(5)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15
        b = d.styles["List Bullet"]
        b.font.name = FONT; b.font.size = Pt(12)
        b.paragraph_format.line_spacing = 1.3; b.paragraph_format.space_after = Pt(2)
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.54)
        s.top_margin, s.bottom_margin = Cm(2.3), Cm(2.3)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp, "PAGE", "1")

    # ------------------------------------------------------------------ text
    def runs(self, p, text, size=12, bold=False, italic=False, color=None):
        for tok in re.split(r"(\[\[.+?\]\]|\*\*.+?\*\*|`.+?`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))", text):
            if not tok: continue
            if tok.startswith("[["):
                _font(p.add_run("[" + tok[2:-2] + "]"), size, True, italic, highlight=True)
            elif tok.startswith("**"):
                _font(p.add_run(tok[2:-2]), size, True, italic, color)
            elif tok.startswith("`"):
                _font(p.add_run(tok[1:-1]), size - 1, bold, italic, name="Courier New")
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                _font(p.add_run(tok[1:-1]), size, bold, True, color)
            else:
                _font(p.add_run(tok), size, bold, italic, color)

    def para(self, text, align="j", size=12, bold=False, italic=False, color=None, before=0, after=6, spacing=1.3,
             indent=None, keep=False):
        p = self.doc.add_paragraph()
        p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                       "l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
        if indent is not None: pf.left_indent = Cm(indent)
        if keep: pf.keep_with_next = True
        self.runs(p, text, size, bold, italic, color)
        return p

    def h(self, level, text):
        p = self.doc.add_heading(level=level)
        _font(p.add_run(text), {1: 14, 2: 12.5, 3: 12}[level], True, False, NAVY)
        return p

    def bullets(self, items, align="j"):
        for t in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align == "j" else WD_ALIGN_PARAGRAPH.LEFT
            self.runs(p, t)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def caption(self, text):
        return self.para(text, align="c", size=9.5, italic=True, color=CAPBLUE, after=6, spacing=1.0, keep=False)

    # ----------------------------------------------------------------- tables
    def table(self, title, header, rows, widths, size=10.5, center=True):
        n = len(self.tables) + 1
        cap = f"Table {n} - {title}"
        self.tables.append(cap)
        p = self.caption(cap); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""
            q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.line_spacing = 1.0; q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3)
            self.runs(q, h, size, True)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].text = ""
                q = cells[i].paragraphs[0]
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
                q.paragraph_format.line_spacing = 1.0; q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3)
                self.runs(q, str(txt), size)
        for r in t.rows:
            c = OxmlElement("w:cantSplit"); c.set(qn("w:val"), "true"); r._tr.get_or_add_trPr().append(c)
        fix_widths(t, widths)
        if len(rows) <= 12:
            for row in t.rows[:-1]:
                for c in row.cells:
                    for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
        self.para("", after=4, spacing=1.0)
        return t

    # ---------------------------------------------------------------- figures
    def figure(self, path, title, width_cm=14.0):
        n = len(self.figures) + 1
        cap = f"Figure {n} - {title}"
        self.figures.append(cap)
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.caption(cap)

    def placeholder_box(self, text, title, height_cm=5.0):
        n = len(self.figures) + 1
        cap = f"Figure {n} - {title}"
        self.figures.append(cap)
        t = self.doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; c.text = ""
        q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = Pt(height_cm * 12); q.paragraph_format.space_after = Pt(height_cm * 12)
        self.runs(q, text, 11)
        fix_widths(t, [13.0])
        self.caption(cap)

    # --------------------------------------------------------------- sections
    def landscape(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.6); s.top_margin, s.bottom_margin = Cm(1.4), Cm(1.8)
        return s

    def portrait(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.54); s.top_margin, s.bottom_margin = Cm(2.3), Cm(2.3)
        return s

    def save(self, path):
        self.doc.save(path)
